#!/usr/bin/env python3
"""
extract_xlsx_to_jsonl.py
========================
Generic xlsx → JSONL extractor.

Reads an xlsx file and writes each data row as a JSON object.
Column mapping is specified via --columns flag: col_index:field_name pairs.

Usage:
    python extract_xlsx_to_jsonl.py \
        --input "datasets/hadith-data/sheikahmad/Hadith-Pure-Canon-Authentica.xlsx" \
        --output datasets/hadith-data/pure_canon.jsonl \
        --sheet Sheet1 \
        --columns 0:row_no 1:text_ar 2:ruling 3:source 6:topics 7:chain 12:shelf_no 13:hadith_no

    python extract_xlsx_to_jsonl.py \
        --input "datasets/hadith-data/sheikahmad/NabiMuhammad-com_First_Green_Book_of_Sahih_Hadith.xlsx" \
        --output datasets/hadith-data/green_book.jsonl \
        --sheet Sheet1 \
        --columns 0:sn 2:match_pct 5:source_book 6:full_text 7:text_ar
"""

import argparse
import io
import json
import sys

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")


def parse_columns(col_spec: list[str]) -> list[tuple[int, str]]:
    """Parse '0:row_no 1:text_ar' into [(0, 'row_no'), (1, 'text_ar')]."""
    result = []
    for spec in col_spec:
        idx_str, name = spec.split(":", 1)
        result.append((int(idx_str), name))
    return result


def main():
    parser = argparse.ArgumentParser(description="Extract xlsx → JSONL")
    parser.add_argument("--input", required=True, help="Input xlsx path")
    parser.add_argument("--output", required=True, help="Output JSONL path")
    parser.add_argument("--sheet", default="Sheet1", help="Sheet name")
    parser.add_argument(
        "--columns",
        nargs="+",
        required=True,
        help="Column mappings: col_index:field_name (e.g. 0:row_no 1:text_ar)",
    )
    parser.add_argument(
        "--id-template",
        default=None,
        help="Template for id field using {field_name} placeholders (e.g. '{source}_{hadith_no}')",
    )
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    import os

    if os.path.exists(args.output) and not args.overwrite:
        print(f"ERROR: {args.output} exists. Use --overwrite to replace.")
        sys.exit(1)

    col_map = parse_columns(args.columns)
    print(f"Reading: {args.input} (sheet: {args.sheet})")
    print(f"Columns: {col_map}")

    wb = openpyxl.load_workbook(args.input, read_only=True)
    ws = wb[args.sheet]

    count = 0
    with open(args.output, "w", encoding="utf-8") as out:
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
            obj = {}
            for col_idx, field_name in col_map:
                val = row[col_idx] if col_idx < len(row) else None
                obj[field_name] = str(val).strip() if val is not None else ""

            # Generate id if template provided
            if args.id_template:
                try:
                    obj["id"] = args.id_template.format(**obj)
                except KeyError:
                    obj["id"] = f"row_{row_idx}"
            elif "id" not in obj:
                obj["id"] = f"row_{row_idx}"

            # Skip rows with no text_ar (empty source text)
            if "text_ar" in obj and not obj["text_ar"]:
                continue

            out.write(json.dumps(obj, ensure_ascii=False) + "\n")
            count += 1

            if count % 10000 == 0:
                print(f"  Extracted {count} rows...")

    wb.close()
    print(f"Done. {count} rows written to {args.output}")


if __name__ == "__main__":
    main()
